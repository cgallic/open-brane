#!/usr/bin/env python3
"""
gdrive_extractor.py — extract text from Google Drive files into a normalized
JSONL stream for brain ingestion.

Runs on the machine that has Google Drive for Desktop mounted (typically your
laptop). Output is one JSON object per line (one per chunk), printed to stdout
or --output path. Designed to be scp'd to the VPS and loaded via
ingest_gdrive.py.

Two extraction paths:
  1. Local files (PDF, DOCX, CSV, TXT, etc.) — read directly from the mount.
  2. Google-native files (.gdoc, .gsheet) — use the Drive API to export as
     text/CSV, matched by filename stem against a bulk file listing.

Chunking:
  Files larger than chunk_size_chars are split into overlapping chunks. Each
  chunk becomes one JSONL record. Boundary refinement scans backward for
  paragraph, sentence, then word breaks.

Google API setup (only needed for .gdoc/.gsheet):
  1. Console: create a project, enable Google Drive API, create OAuth 2.0
     Client ID (type: Desktop app). Download JSON.
  2. Save it to config/gdrive-credentials.json.
  3. First run will open a browser for consent; token saved to
     config/gdrive-token.json and refreshed automatically.

Usage:
    gdrive_extractor.py [--drive-root "G:/My Drive"] [--config PATH]
                        [--output PATH] [--limit N] [--verbose]
                        [--dry-run] [--no-google-api]

Dependencies (install only what you need):
    pip install pdfplumber python-docx openpyxl python-pptx
    pip install google-auth google-auth-oauthlib google-api-python-client
"""
import argparse
import csv
import email
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
CONFIG_DEFAULT = SCRIPT_DIR.parent / "config" / "gdrive-sources.json"

DEFAULT_DRIVE_ROOT = "G:/My Drive"
DEFAULT_INCLUDE_EXT = [".pdf", ".docx", ".txt", ".csv", ".md", ".xlsx", ".xls",
                       ".pptx", ".html", ".htm", ".json", ".eml", ".tsv", ".rtf"]
DEFAULT_GOOGLE_EXT = [".gdoc", ".gsheet"]
DEFAULT_EXCLUDE_FOLDERS: list[str] = []
DEFAULT_MAX_FILE_SIZE_MB = 50
DEFAULT_CHUNK_SIZE = 3800
DEFAULT_CHUNK_OVERLAP = 200
DEFAULT_MAX_CHUNKS = 20
HASH_PREFIX_CHARS = 16


# --- Google Drive API (lazy-loaded) ---------------------------------------

_gdrive_service = None


def _get_gdrive_service(credentials_path: str, token_path: str):
    global _gdrive_service
    if _gdrive_service is not None:
        return _gdrive_service

    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
    creds = None

    token_p = Path(token_path)
    creds_p = Path(credentials_path)

    if token_p.exists():
        creds = Credentials.from_authorized_user_file(str(token_p), SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not creds_p.exists():
                raise FileNotFoundError(
                    f"Google API credentials not found: {creds_p}\n"
                    "Download OAuth 2.0 Client ID JSON from Google Cloud Console "
                    "and save it there. Or use --no-google-api to skip."
                )
            flow = InstalledAppFlow.from_client_secrets_file(str(creds_p), SCOPES)
            creds = flow.run_local_server(port=0)
        token_p.parent.mkdir(parents=True, exist_ok=True)
        token_p.write_text(creds.to_json())

    _gdrive_service = build("drive", "v3", credentials=creds)
    return _gdrive_service


_gdrive_file_index: dict[str, list[dict]] | None = None


def _build_gdrive_index(service, verbose: bool = False) -> dict[str, list[dict]]:
    """Bulk-list all Google Docs and Sheets. Drive-for-Desktop .gdoc/.gsheet
    stubs are unreadable, so we list via API and match by filename stem."""
    global _gdrive_file_index
    if _gdrive_file_index is not None:
        return _gdrive_file_index

    index: dict[str, list[dict]] = {}
    mime_types = [
        "application/vnd.google-apps.document",
        "application/vnd.google-apps.spreadsheet",
    ]
    query = " or ".join(f"mimeType='{m}'" for m in mime_types)
    query = f"({query}) and trashed=false"
    fields = "nextPageToken, files(id, name, mimeType, modifiedTime)"

    page_token = None
    total = 0
    while True:
        resp = service.files().list(
            q=query, fields=fields, pageSize=1000, pageToken=page_token,
        ).execute()
        for f in resp.get("files", []):
            stem = f["name"].lower()
            index.setdefault(stem, []).append(f)
            total += 1
        page_token = resp.get("nextPageToken")
        if not page_token:
            break

    if verbose:
        print(f"[google-api] indexed {total} Google Docs/Sheets", file=sys.stderr)

    _gdrive_file_index = index
    return index


def _find_google_file(stem: str, mime_type: str, index: dict[str, list[dict]]) -> str | None:
    candidates = index.get(stem.lower(), [])
    for c in candidates:
        if c["mimeType"] == mime_type:
            return c["id"]
    return None


def _decode_export(raw) -> str | None:
    if isinstance(raw, bytes):
        text = raw.decode("utf-8", errors="replace")
    else:
        text = str(raw)
    text = text.lstrip("﻿").replace("\r\n", "\n").strip()
    return text if text else None


def extract_google_doc(path: Path, service, index) -> str | None:
    file_id = _find_google_file(path.stem, "application/vnd.google-apps.document", index)
    if not file_id:
        return None
    try:
        raw = service.files().export(fileId=file_id, mimeType="text/plain").execute()
        return _decode_export(raw)
    except Exception:
        return None


def extract_google_sheet(path: Path, service, index) -> str | None:
    file_id = _find_google_file(path.stem, "application/vnd.google-apps.spreadsheet", index)
    if not file_id:
        return None
    try:
        raw = service.files().export(fileId=file_id, mimeType="text/csv").execute()
        return _decode_export(raw)
    except Exception:
        return None


# --- Local file extractors ------------------------------------------------

def extract_pdf(path: Path) -> str | None:
    import pdfplumber
    try:
        pages = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
        result = "\n\n".join(pages)
        return result if len(result) >= 50 else None
    except Exception:
        return None


def extract_docx(path: Path) -> str | None:
    from docx import Document
    try:
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        result = "\n\n".join(paragraphs)
        return result if result.strip() else None
    except Exception:
        return None


def extract_xlsx(path: Path) -> str | None:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        sheet = wb.active
        if sheet is None:
            return None
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= 200:
                break
            cells = [str(c) if c is not None else "" for c in row]
            rows.append(",".join(cells))
        wb.close()
        result = "\n".join(rows)
        return result if result.strip() else None
    except Exception:
        return None


def extract_pptx(path: Path) -> str | None:
    from pptx import Presentation
    try:
        prs = Presentation(str(path))
        slides_text = []
        for i, slide in enumerate(prs.slides):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            parts.append(t)
            if parts:
                slides_text.append(f"[Slide {i+1}]\n" + "\n".join(parts))
        result = "\n\n".join(slides_text)
        return result if result.strip() else None
    except Exception:
        return None


def extract_csv_file(path: Path) -> str | None:
    try:
        with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(f, dialect)
            rows = []
            for i, row in enumerate(reader):
                if i >= 200:
                    break
                rows.append(",".join(row))
        result = "\n".join(rows)
        return result if result.strip() else None
    except Exception:
        return None


def extract_text(path: Path) -> str | None:
    for enc in ["utf-8", "utf-16", "latin-1"]:
        try:
            text = path.read_text(encoding=enc, errors="strict")
            return text if text.strip() else None
        except (UnicodeDecodeError, UnicodeError):
            continue
        except OSError:
            return None
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None


def extract_html(path: Path) -> str | None:
    text = extract_text(path)
    if not text:
        return None
    text = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) >= 50 else None


def extract_json_file(path: Path) -> str | None:
    try:
        if path.stat().st_size > 1_000_000:
            return None
        text = path.read_text(encoding="utf-8", errors="replace")
        data = json.loads(text)
        return json.dumps(data, indent=2, ensure_ascii=False)
    except (json.JSONDecodeError, OSError):
        return None


def extract_email_file(path: Path) -> str | None:
    try:
        raw = path.read_bytes()
        msg = email.message_from_bytes(raw)
        parts = [
            f"From: {msg.get('From', '')}",
            f"To: {msg.get('To', '')}",
            f"Subject: {msg.get('Subject', '')}",
            f"Date: {msg.get('Date', '')}",
            "",
        ]
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    payload = part.get_payload(decode=True)
                    if payload:
                        parts.append(payload.decode("utf-8", errors="replace"))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                parts.append(payload.decode("utf-8", errors="replace"))
        result = "\n".join(parts)
        return result if len(result) > 50 else None
    except Exception:
        return None


LOCAL_HANDLERS: dict = {
    ".pdf": extract_pdf, ".docx": extract_docx, ".xlsx": extract_xlsx,
    ".xls": extract_xlsx, ".pptx": extract_pptx, ".csv": extract_csv_file,
    ".tsv": extract_csv_file, ".txt": extract_text, ".md": extract_text,
    ".rtf": extract_text, ".html": extract_html, ".htm": extract_html,
    ".json": extract_json_file, ".eml": extract_email_file,
}


def chunk_text(text: str, chunk_size: int, overlap: int, max_chunks: int) -> list[str]:
    if len(text) <= chunk_size:
        return [text]
    chunks = []
    start = 0
    while start < len(text) and len(chunks) < max_chunks:
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:])
            break
        candidate = end
        para_pos = text.rfind("\n\n", start + chunk_size - 200, end)
        if para_pos > start:
            candidate = para_pos + 2
        else:
            sent_pos = text.rfind(". ", start + chunk_size - 200, end)
            if sent_pos > start:
                candidate = sent_pos + 2
            else:
                word_pos = text.rfind(" ", start + chunk_size - 200, end)
                if word_pos > start:
                    candidate = word_pos + 1
        chunks.append(text[start:candidate])
        start = max(0, candidate - overlap)
    if len(chunks) >= max_chunks and start < len(text):
        tail = text[-chunk_size:]
        if tail != chunks[-1]:
            chunks[-1] = tail
    return chunks


def make_hash(relative_path: str, chunk_idx: int) -> str:
    key = f"{relative_path}:{chunk_idx}".encode("utf-8", errors="replace")
    return hashlib.sha256(key).hexdigest()[:HASH_PREFIX_CHARS]


def file_mtime_iso(path: Path) -> str:
    try:
        return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()
    except OSError:
        return datetime.now(tz=timezone.utc).isoformat()


def load_config(config_path: Path | None) -> dict:
    if config_path and config_path.exists():
        return json.loads(config_path.read_text(encoding="utf-8"))
    if CONFIG_DEFAULT.exists():
        return json.loads(CONFIG_DEFAULT.read_text(encoding="utf-8"))
    return {}


def main():
    ap = argparse.ArgumentParser(description="Extract Google Drive files to JSONL for ingestion.")
    ap.add_argument("--drive-root", default=None)
    ap.add_argument("--config", default=None)
    ap.add_argument("--output", default="-", help="output path (default: stdout)")
    ap.add_argument("--limit", type=int, default=0, help="max files (0=unlimited)")
    ap.add_argument("--verbose", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-google-api", action="store_true")
    args = ap.parse_args()

    cfg = load_config(Path(args.config) if args.config else None)

    drive_root = Path(args.drive_root or cfg.get("drive_root", DEFAULT_DRIVE_ROOT))
    include_ext = set(cfg.get("include_extensions", DEFAULT_INCLUDE_EXT))
    google_ext = set(cfg.get("google_extensions", DEFAULT_GOOGLE_EXT))
    exclude_folders = cfg.get("exclude_folders", DEFAULT_EXCLUDE_FOLDERS)
    max_file_size = cfg.get("max_file_size_mb", DEFAULT_MAX_FILE_SIZE_MB) * 1024 * 1024
    chunk_size = cfg.get("chunk_size_chars", DEFAULT_CHUNK_SIZE)
    chunk_overlap = cfg.get("chunk_overlap_chars", DEFAULT_CHUNK_OVERLAP)
    max_chunks = cfg.get("max_chunks_per_file", DEFAULT_MAX_CHUNKS)

    pkg_root = SCRIPT_DIR.parent
    creds_path = str(pkg_root / cfg.get("google_api_credentials", "config/gdrive-credentials.json"))
    token_path = str(pkg_root / cfg.get("google_api_token", "config/gdrive-token.json"))

    if not drive_root.exists():
        print(json.dumps({"status": "error", "error": f"drive root missing: {drive_root}"}), file=sys.stderr)
        sys.exit(1)

    all_ext = include_ext | (set() if args.no_google_api else google_ext)

    gdrive_service = None
    gdrive_index: dict = {}
    google_api_available = not args.no_google_api
    google_api_errors = 0

    files_found: list[tuple[Path, str]] = []
    skipped_excluded = skipped_ext = skipped_size = skipped_google = 0

    if args.verbose:
        print(f"[extract] walking {drive_root}...", file=sys.stderr)

    for dirpath, dirnames, filenames in os.walk(str(drive_root)):
        dp = Path(dirpath)
        dirnames[:] = [d for d in dirnames if d.lower() not in {f.lower() for f in exclude_folders}]
        for fname in filenames:
            fpath = dp / fname
            ext = fpath.suffix.lower()
            if ext not in all_ext:
                skipped_ext += 1
                continue
            if ext in google_ext and args.no_google_api:
                skipped_google += 1
                continue
            try:
                fsize = fpath.stat().st_size
            except OSError:
                continue
            if fsize > max_file_size:
                skipped_size += 1
                continue
            files_found.append((fpath, ext))

    if args.verbose:
        print(f"[extract] {len(files_found)} files found", file=sys.stderr)

    files_found.sort(key=lambda x: x[0].stat().st_mtime)
    if args.limit > 0:
        files_found = files_found[:args.limit]

    out_fp = sys.stdout if args.output == "-" else open(args.output, "w", encoding="utf-8")
    total_chunks = files_ok = files_error = files_empty = 0

    try:
        for fpath, ext in files_found:
            try:
                rel_path = str(fpath.relative_to(drive_root)).replace("\\", "/")
            except ValueError:
                rel_path = fpath.name
            folder = str(Path(rel_path).parent) if "/" in rel_path else ""

            if args.dry_run:
                files_ok += 1
                continue

            text = None
            if ext in google_ext:
                if gdrive_service is None and google_api_available:
                    try:
                        gdrive_service = _get_gdrive_service(creds_path, token_path)
                        gdrive_index = _build_gdrive_index(gdrive_service, verbose=args.verbose)
                    except Exception as e:
                        if args.verbose:
                            print(f"[error] Google API init failed: {e}", file=sys.stderr)
                        google_api_available = False
                        google_api_errors += 1
                if gdrive_service:
                    if ext == ".gdoc":
                        text = extract_google_doc(fpath, gdrive_service, gdrive_index)
                    elif ext == ".gsheet":
                        text = extract_google_sheet(fpath, gdrive_service, gdrive_index)
                    if text is None:
                        google_api_errors += 1
            else:
                handler = LOCAL_HANDLERS.get(ext)
                if handler:
                    text = handler(fpath)

            if not text or not text.strip():
                files_empty += 1
                continue

            chunks = chunk_text(text, chunk_size, chunk_overlap, max_chunks)
            try:
                fsize = fpath.stat().st_size
            except OSError:
                fsize = 0
            mtime = file_mtime_iso(fpath)

            for ci, chunk_content in enumerate(chunks):
                record = {
                    "file_path": rel_path,
                    "file_name": fpath.name,
                    "file_ext": ext,
                    "file_size": fsize,
                    "file_mtime_iso": mtime,
                    "folder": folder,
                    "content": chunk_content,
                    "chunk_idx": ci,
                    "chunk_total": len(chunks),
                    "word_count": len(chunk_content.split()),
                    "hash": make_hash(rel_path, ci),
                }
                out_fp.write(json.dumps(record, ensure_ascii=False) + "\n")
                total_chunks += 1
            files_ok += 1

    except KeyboardInterrupt:
        print("\n[interrupted]", file=sys.stderr)
    except Exception as e:
        files_error += 1
        print(f"[error] {e}", file=sys.stderr)
    finally:
        if out_fp is not sys.stdout:
            out_fp.close()

    print(json.dumps({
        "status": "ok", "files_found": len(files_found), "files_ok": files_ok,
        "files_empty": files_empty, "files_error": files_error,
        "chunks_extracted": total_chunks, "skipped_wrong_ext": skipped_ext,
        "skipped_excluded": skipped_excluded, "skipped_too_large": skipped_size,
        "skipped_google_no_api": skipped_google, "google_api_errors": google_api_errors,
        "dry_run": args.dry_run,
    }), file=sys.stderr)


if __name__ == "__main__":
    main()
